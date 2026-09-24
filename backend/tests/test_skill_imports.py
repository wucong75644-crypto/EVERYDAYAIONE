"""Real document import, bounded parsing, original hashes and safe API access."""
import base64
import hashlib
from io import BytesIO
import multiprocessing
from unittest.mock import Mock
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from openpyxl import Workbook
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest

from services.skills import imports
from services.skills.assets import AssetSourceDraft, SkillResources, public_assets
from services.skills.authoring_contracts import DraftContent, reviewed_document
from services.skills.contracts import SkillError
from services.skills.renderer import asset_manifest_digest, digest, encoded, prepare_resources, MAX_RENDERED_BYTES
from tests.test_skill_assets import attachment, document, publish
from tests.test_skill_storage import storage, PACKAGE  # noqa: F401
from tests.test_skill_authoring_api import api  # noqa: F401


def office_file(format, value='Imported evidence'):
    output = BytesIO()
    if format == 'docx':
        doc = Document()
        doc.add_paragraph(value)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Item'
        table.cell(0, 1).text = 'Amount'
        doc.save(output)
    elif format == 'xlsx':
        book = Workbook()
        book.active.append([value, 12, '=1+2'])
        book.save(output)
    else:
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 10 100 Td ({value}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize('format', ['docx', 'xlsx', 'pdf', 'txt', 'md', 'csv', 'json'])
def test_upload_reads_real_files_and_keeps_original_bytes(format):
    raw = office_file(format) if format in ('docx', 'xlsx', 'pdf') else b'\xef\xbb\xbfImported evidence\r\n'
    asset = imports.import_attachment(f'Reference.{format}', raw)
    assert asset.name == f'Reference.{format}' and asset.kind == 'reference'
    assert 'Imported evidence' in asset.content
    assert asset.source.raw() == raw
    assert asset.manifest().source.sha256 == hashlib.sha256(raw).hexdigest()
    if format == 'xlsx':
        assert '=1+2' in asset.content  # Formula is inert text, never evaluated.
    if format == 'docx':
        assert 'Item\tAmount' in asset.content


@pytest.mark.parametrize('filename,raw,code', [
    ('../secret.txt', b'a', 'FILENAME_INVALID'), ('C:\\secret.txt', b'a', 'FILENAME_INVALID'),
    ('script.py', b'print(1)', 'FORMAT'), ('legacy.doc', b'a', 'FORMAT'), ('image.png', b'a', 'FORMAT'),
    ('a.txt', b'', 'NO_TEXT'), ('a.txt', b'\xff', 'ENCODING'), ('a.txt', b'\x00text', 'INVALID'),
    ('a.txt', b'x' * 65537, 'TEXT_TOO_LARGE'), ('a.pdf', b'x' * (2 * 1024 * 1024 + 1), 'TOO_LARGE'),
    ('a.pdf', b'not a pdf', 'INVALID'),
])
def test_unsupported_and_unreadable_uploads_fail_before_draft_mutation(filename, raw, code):
    with pytest.raises(SkillError, match=code):
        imports.import_attachment(filename, raw)


def test_blank_and_encrypted_pdf_are_actionable():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    raw = BytesIO()
    writer.write(raw)
    with pytest.raises(SkillError, match='NO_TEXT'):
        imports.import_attachment('scan.pdf', raw.getvalue())
    writer.encrypt('password')
    raw = BytesIO()
    writer.write(raw)
    with pytest.raises(SkillError, match='ENCRYPTED'):
        imports.import_attachment('locked.pdf', raw.getvalue())


@pytest.mark.parametrize('entry,value,code', [
    ('../escape.xml', b'x', 'INVALID'), ('word/vbaProject.bin', b'x', 'INVALID'),
    ('word/bomb.xml', b'x' * (8 * 1024 * 1024 + 1), 'COMPLEXITY_LIMIT'),
    ('word/entity.xml', b'<!DOCTYPE xml [<!ENTITY x "abc">]><xml/>', 'INVALID'),
])
def test_office_container_rejects_paths_macros_entities_and_expansion(entry, value, code):
    raw = BytesIO(office_file('docx'))
    with ZipFile(raw, 'a', compression=ZIP_DEFLATED) as archive:
        archive.writestr(entry, value)
    with pytest.raises(SkillError, match=code):
        imports.import_attachment('unsafe.docx', raw.getvalue())


@pytest.mark.parametrize('cell', ['CW1', 'A2001'])
def test_spreadsheet_limits_do_not_silently_drop_faraway_cells(cell):
    book = Workbook()
    book.active['A1'] = 'Visible'
    book.active[cell] = 'Must not be silently omitted'
    raw = BytesIO()
    book.save(raw)
    with pytest.raises(SkillError, match='COMPLEXITY_LIMIT'):
        imports.import_attachment('large.xlsx', raw.getvalue())


def test_parser_timeout_terminates_worker_and_releases_capacity(monkeypatch):
    before = {p.pid for p in multiprocessing.active_children()}
    monkeypatch.setattr(imports, 'PARSE_TIMEOUT', 0)
    with pytest.raises(SkillError, match='COMPLEXITY_LIMIT'):
        imports.import_attachment('slow.pdf', office_file('pdf'))
    assert {p.pid for p in multiprocessing.active_children()} == before
    assert imports._PARSERS.acquire(blocking=False)
    assert imports._PARSERS.acquire(blocking=False)
    try:
        with pytest.raises(SkillError, match='BUSY'):
            imports.import_attachment('busy.pdf', office_file('pdf'))
    finally:
        imports._PARSERS.release()
        imports._PARSERS.release()


def test_admin_upload_uses_target_org_and_does_not_write_workspace(api):
    response = api.client.post(api.base + '/attachments/import', headers=api.auth,
        files={'file': ('guide.txt', b'Read this')})
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    result = response.json()
    assert result['name'] == 'guide.txt' and result['content'] == 'Read this'
    assert base64.b64decode(result['source']['base64']) == b'Read this'
    assert not {'path', 'url', 'sha256'} & result.keys()
    api.service.create.assert_not_called()
    api.service.save.assert_not_called()
    assert api.client.post(api.base + '/attachments/import', files={'file': ('a.txt', b'a')}).status_code == 401
    assert api.client.post(f'/api/skills/admin/orgs/{api.foreign}/attachments/import', headers=api.auth,
        files={'file': ('a.txt', b'a')}).status_code == 403
    api.db._tables['org_members']._data[0]['role'] = 'member'
    assert api.client.post(api.base + '/attachments/import', headers=api.auth,
        files={'file': ('a.txt', b'a')}).status_code == 403


def test_upload_error_code_is_safe_and_decodable(api):
    response = api.client.post(api.base + '/attachments/import', headers=api.auth,
        files={'file': ('script.py', b'private content')})
    assert response.status_code == 422
    assert response.json()['error']['code'] == 'SKILL_UPLOAD_FORMAT'
    assert 'private content' not in response.text


def uploaded(content=b'Original file'):
    return attachment(content=content.decode(), source=AssetSourceDraft.from_bytes('txt', content))


def test_sources_are_revision_owned_hashed_and_summary_only_until_referenced(storage):
    skill = publish(storage, document(uploaded(), body='Summary only'))
    original = (storage.root / skill.nas_path).parent / 'assets/guide.original.txt'
    assert original.read_bytes() == b'Original file' and original.stat().st_mode & 0o222 == 0
    summary = public_assets(skill.resources)[0]
    assert summary['file_format'] == 'txt' and summary['file_bytes'] == 13
    assert not {'path', 'source', 'sha256', 'base64'} & summary.keys()
    storage._read = Mock(wraps=storage._read)
    ids, _, _ = prepare_resources(skill, {}, MAX_RENDERED_BYTES)
    assert ids == ()
    storage.read_assets(skill, ids)
    storage._read.assert_not_called()
    assert storage.read_assets(skill, ['guide']) == {'guide': 'Original file'}
    assert storage._read.call_count == 2
    assert storage.read_source(skill, 'guide').raw() == b'Original file'


@pytest.mark.parametrize('damage', ['drift', 'symlink', 'missing'])
def test_original_source_damage_fails_closed(storage, tmp_path, damage):
    skill = publish(storage, document(uploaded()))
    original = (storage.root / skill.nas_path).parent / 'assets/guide.original.txt'
    original.unlink()
    if damage == 'drift':
        original.write_bytes(b'Different file')
    elif damage == 'symlink':
        other = tmp_path / 'outside.txt'
        other.write_bytes(b'Original file')
        original.symlink_to(other)
    with pytest.raises(SkillError):
        storage.read_assets(skill, ['guide'])


@pytest.mark.parametrize('path', ['../original.txt', '/tmp/original.txt', 'assets/guide.original.py', 'assets/other.original.txt'])
def test_source_manifest_rejects_paths(path):
    manifest = uploaded().manifest().model_dump()
    manifest['source']['path'] = path
    with pytest.raises(ValidationError):
        SkillResources(assets=[manifest])


def test_original_bytes_are_required_and_covered_by_review_hash(storage):
    content = document(uploaded())
    publication, raw = reviewed_document(PACKAGE, 'v1', content)
    for sources in ({}, {'guide': b'tampered'}):
        with pytest.raises(SkillError):
            storage.publish(PACKAGE, publication, raw, assets={'guide': b'Original file'}, sources=sources)
    changed = content.model_copy(update={'assets': (attachment(content='Original file',
        source=AssetSourceDraft.from_bytes('txt', b'Changed file')),)})
    assert reviewed_document(PACKAGE, 'v1', changed)[0].content_sha256 != publication.content_sha256
    with pytest.raises(ValidationError):
        DraftContent(assets=tuple(attachment(id=f'file-{i}', source=AssetSourceDraft.from_bytes('txt', b'x' * 2_000_000)) for i in range(5)))
    with pytest.raises(ValidationError):
        AssetSourceDraft(format='pdf', base64='!!!!')


def test_legacy_manifest_and_review_hash_are_unchanged():
    assert 'source' not in document(attachment()).model_dump(mode='json')['assets'][0]
    resources = SkillResources(assets=[attachment().manifest()])
    old = resources.model_dump(mode='json')
    old['assets'][0].pop('source')
    assert asset_manifest_digest(resources) == digest(encoded(old))
    _, raw = reviewed_document(PACKAGE, 'v1', document(attachment()))
    assert b'source:' not in raw
