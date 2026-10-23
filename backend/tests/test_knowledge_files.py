"""Private file HTTP, lifecycle, canonical registration and immutable references."""
from pathlib import Path
from uuid import uuid4
from types import SimpleNamespace
import psycopg
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from tests.test_knowledge_workspace_v2 import local_db, command, document
from services.knowledge_workspace.repository import Repository

@pytest.fixture
def file_client(local_db,tmp_path,monkeypatch):
    from api.deps import get_org_context, OrgContext
    from core.database import get_db
    from api.routes import knowledge, knowledge_files
    make,actors=local_db;r=make();ctx=OrgContext(actors[0],r.scope.org_id)
    settings=SimpleNamespace(knowledge_workspace_enabled=True,knowledge_workspace_user_ids='',knowledge_workspace_org_ids='',file_workspace_enabled=True,file_workspace_root=str(tmp_path/'nas'))
    monkeypatch.setattr(knowledge,'get_settings',lambda:settings);monkeypatch.setattr(knowledge_files,'get_settings',lambda:settings)
    app=FastAPI();app.include_router(knowledge.router,prefix='/api');app.include_router(knowledge_files.router,prefix='/api')
    app.dependency_overrides[get_org_context]=lambda:ctx
    app.dependency_overrides[get_db]=lambda:SimpleNamespace(pool=r.pool)
    with TestClient(app,raise_server_exceptions=False) as client:yield client,ctx,settings,r,make

def uploaded(client,oid,body='附件正文 中文验收',name='资料.txt',key=None,**data):
    response=client.post(f'/api/knowledge/objects/{oid}/files',data={'command_id':str(key or uuid4()),**data},files={'file':(name,body.encode() if isinstance(body,str) else body,'application/octet-stream')})
    assert response.status_code==200,response.text
    return response.json()['data']

def url(oid,bid,rev=1):return f'/api/knowledge/objects/{oid}/files/{bid}/versions/{rev}/content'

def test_file_versions_references_removal_and_acl(file_client,local_db):
    client,ctx,_,r,make=file_client;_,actors=local_db;original=document(r);target=document(r);oid=original['id'];tid=target['id']
    bid=uploaded(client,oid)['binding_id']
    listing=client.get(f'/api/knowledge/objects/{oid}/files');assert listing.status_code==200
    metadata=listing.json()['data'][0]['version'];assert metadata['sha256'] and 'storage_owner' not in metadata and 'storage_org' not in metadata
    assert 'url' not in str(listing.json()) and 'workspace_path' not in str(listing.json())
    result=client.get(url(oid,bid));assert result.content=='附件正文 中文验收'.encode()
    assert result.headers['cache-control']=='no-store' and result.headers['x-content-type-options']=='nosniff'
    assert r.search(original['container_id'],'中文验收')['items']
    copied=client.post(f'/api/knowledge/objects/{tid}/files/reference',json={'command_id':str(uuid4()),'source_object_id':str(oid),'source_binding_id':bid,'source_revision':1})
    assert copied.status_code==200,copied.text;copy_id=copied.json()['data']['binding_id']
    uploaded(client,oid,'新的附件内容',binding_id=bid,expected_revision='1')
    assert client.get(url(oid,bid)).content=='附件正文 中文验收'.encode()
    assert client.get(url(oid,bid,2)).content=='新的附件内容'.encode()
    assert client.get(url(tid,copy_id)).content=='附件正文 中文验收'.encode()
    assert not r.search(original['container_id'],'中文验收')['items'] and r.search(target['container_id'],'中文验收')['items']
    versions=client.get(f'/api/knowledge/objects/{oid}/files/{bid}/versions').json()['data']['items'];assert [v['revision'] for v in versions]==[2,1]
    r.command(command('set_member',{'container_id':target['container_id'],'user_id':actors[1],'role':'viewer','status':'active'}));ctx.user_id=actors[1]
    assert client.get(url(tid,copy_id)).status_code==200 and client.get(url(oid,bid)).status_code==404
    assert client.get(url(tid,bid)).status_code==404
    denied=client.post(f'/api/knowledge/objects/{tid}/files',data={'command_id':str(uuid4())},files={'file':('禁止.txt',b'no')});assert denied.status_code==403
    ctx.org_id=make(1,1).scope.org_id;assert client.get(url(tid,copy_id)).status_code==404;ctx.org_id=r.scope.org_id
    r.command(command('set_member',{'container_id':target['container_id'],'user_id':actors[1],'role':'viewer','status':'removed'}))
    assert client.get(url(tid,copy_id)).status_code==404
    ctx.user_id=actors[0]
    removed=client.request('DELETE',f'/api/knowledge/objects/{oid}/files/{bid}',json={'command_id':str(uuid4()),'expected_revision':2});assert removed.status_code==200
    assert client.get(url(oid,bid)).status_code==404
    assert client.get(f'/api/knowledge/objects/{oid}/files/{bid}/versions').status_code==404
    assert client.get(url(tid,copy_id)).status_code==200

def test_workspace_snapshot_survives_source_changes(file_client):
    from services.file_executor import FileExecutor
    client,ctx,settings,r,_=file_client;obj=document(r);executor=FileExecutor(settings.file_workspace_root,ctx.user_id,ctx.org_id)
    source=Path(executor.workspace_root)/'source.txt';source.write_text('固定的源文件正文')
    response=client.post(f'/api/knowledge/objects/{obj["id"]}/files/from-workspace',json={'command_id':str(uuid4()),'workspace_path':'source.txt'})
    assert response.status_code==200,response.text;bid=response.json()['data']['binding_id']
    source.rename(source.with_name('moved.txt'));source=source.with_name('moved.txt');source.write_text('被覆盖');source.unlink()
    assert client.get(url(obj['id'],bid)).content=='固定的源文件正文'.encode()
    denied=client.post(f'/api/knowledge/objects/{obj["id"]}/files/from-workspace',json={'command_id':str(uuid4()),'workspace_path':'../../outside.txt'});assert denied.status_code==404
    with pytest.raises(PermissionError):executor.resolve_safe_path(str(Path(settings.file_workspace_root)/'.knowledge-versions'))

def test_lost_response_replay_and_transaction_rollback(file_client,monkeypatch):
    from services.knowledge_workspace.files import version_id
    client,ctx,_,r,make=file_client;obj=document(r);oid=obj['id'];key=uuid4();original=Repository.file_command
    def lost(self,*args):
        original(self,*args)
        raise RuntimeError('synthetic response loss after commit')
    monkeypatch.setattr(Repository,'file_command',lost)
    response=client.post(f'/api/knowledge/objects/{oid}/files',data={'command_id':str(key)},files={'file':('重试.txt',b'retry')});assert response.status_code==500
    monkeypatch.setattr(Repository,'file_command',original)
    result=uploaded(client,oid,b'retry','重试.txt',key);assert result['status']=='replayed'
    assert len(client.get(f'/api/knowledge/objects/{oid}/files').json()['data'])==1
    changed=client.post(f'/api/knowledge/objects/{oid}/files',data={'command_id':str(key)},files={'file':('重试.txt',b'different')});assert changed.status_code==409
    vid=version_id(ctx.user_id,ctx.org_id,uuid4());payload={'object_id':str(oid),'version_id':vid,'binding_id':None,'expected_revision':None,'name':'bad.txt','mime_type':'text/plain','size':3,'sha256':'0'*64,'extracted_text':'','text_status':'invalid'}
    with pytest.raises(psycopg.Error):r.file_command(uuid4(),'file_upload',payload)
    with psycopg.connect(**make.admin_args) as c:
        assert c.execute('SELECT count(*) FROM user_assets WHERE storage_key=%s',('knowledge-private/'+vid,)).fetchone()[0]==0
        assert c.execute('SELECT count(*) FROM kb_file_versions WHERE id=%s',(vid,)).fetchone()[0]==0

def test_deletion_protection_and_integrity(file_client):
    from services.knowledge_workspace.files import private_path
    client,ctx,settings,r,make=file_client;obj=document(r);bid=uploaded(client,obj['id'])['binding_id']
    version=client.get(f'/api/knowledge/objects/{obj["id"]}/files').json()['data'][0]['version']
    for statement in ['DELETE FROM user_assets WHERE id=%s',"UPDATE user_assets SET status='deleted' WHERE id=%s"]:
        with psycopg.connect(**make.admin_args) as c:
            with pytest.raises(psycopg.Error,match='KB_ASSET_REFERENCED'):c.execute(statement,(version['asset_id'],))
            c.rollback()
    target=private_path(settings.file_workspace_root,ctx.org_id,ctx.user_id,version['version_id']);target.chmod(0o600);target.write_bytes(b'corrupted')
    response=client.get(url(obj['id'],bid));assert response.status_code==503 and response.json()['error']['code']=='KNOWLEDGE_FILE_INTEGRITY'
    assert b'corrupted' not in response.content

@pytest.mark.parametrize('filename,body,status',[
    ('bad.exe',b'payload',422),('fake.png',b'not an image',422),('empty.txt',b'',413),('invalid.txt',b'\xff\xff',422),('bad.docx',b'not a zip',422),('bad.txt',b'null\x00byte',422),
])
def test_invalid_uploads_never_bind(file_client,filename,body,status):
    client,_,_,r,_=file_client;obj=document(r)
    response=client.post(f'/api/knowledge/objects/{obj["id"]}/files',data={'command_id':str(uuid4())},files={'file':(filename,body)})
    assert response.status_code==status,response.text
    assert client.get(f'/api/knowledge/objects/{obj["id"]}/files').json()['data']==[]


def test_second_editor_new_version_reads_from_its_own_private_storage(file_client,local_db):
    client,ctx,_,r,_=file_client;_,actors=local_db;obj=document(r);bid=uploaded(client,obj['id'],name='old-label.txt')['binding_id']
    r.command(command('set_member',{'container_id':obj['container_id'],'user_id':actors[1],'role':'editor','status':'active'}))
    ctx.user_id=actors[1];uploaded(client,obj['id'],'协作者上传新版',name='new-label.txt',binding_id=bid,expected_revision='1')
    ctx.user_id=actors[0]
    assert client.get(url(obj['id'],bid,2)).content=='协作者上传新版'.encode()
    assert r.search(obj['container_id'],'old-label.txt')['items']
    assert client.get(url(obj['id'],bid,1)).content=='附件正文 中文验收'.encode()
    stale=client.post(f'/api/knowledge/objects/{obj["id"]}/files',data={'command_id':str(uuid4()),'binding_id':bid,'expected_revision':'1'},files={'file':('stale.txt',b'stale')})
    assert stale.status_code==409
    assert client.get(url(obj['id'],bid,2)).content=='协作者上传新版'.encode()


def test_file_versions_paginate_and_runtime_cannot_mutate_tables(file_client):
    from services.knowledge_workspace.files import version_id
    client,ctx,_,r,_=file_client;obj=document(r);bid=uploaded(client,obj['id'])['binding_id']
    for revision in range(1,52):
        key=uuid4();vid=version_id(ctx.user_id,ctx.org_id,key)
        r.file_command(key,'file_upload',{'object_id':str(obj['id']),'binding_id':bid,'expected_revision':revision,'version_id':vid,'name':'metadata.txt','mime_type':'text/plain','size':1,'sha256':'0'*64,'text_status':'unsupported'})
    page=client.get(f'/api/knowledge/objects/{obj["id"]}/files/{bid}/versions').json()['data']
    next_page=client.get(f'/api/knowledge/objects/{obj["id"]}/files/{bid}/versions',params={'before':page['next_cursor']}).json()['data']
    assert [v['revision'] for v in page['items']+next_page['items']]==list(range(52,0,-1))
    assert next_page['next_cursor'] is None
    with pytest.raises(psycopg.errors.InsufficientPrivilege):
        with r.cursor() as c:c.execute('SELECT * FROM kb_file_versions')
    with pytest.raises(psycopg.errors.InsufficientPrivilege):
        with r.cursor() as c:c.execute('DELETE FROM kb_file_bindings')


def test_bounded_office_extraction_and_unsupported_pdf(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from services.knowledge_workspace.files import inspect_file
    doc=Document();doc.add_paragraph('Word 的中文正文');path=tmp_path/'word.docx';doc.save(path)
    parsed=inspect_file(path,'word.docx');assert parsed['text_status']=='ready' and '中文正文' in parsed['extracted_text']
    book=Workbook();book.active['A1']='表格中的测试内容';path=tmp_path/'table.xlsx';book.save(path)
    parsed=inspect_file(path,'table.xlsx');assert parsed['text_status']=='ready' and '测试内容' in parsed['extracted_text']
    book.active.cell(row=1,column=201,value='超出索引列数');book.save(path)
    parsed=inspect_file(path,'table.xlsx');assert parsed['text_status']=='truncated' and '超出索引列数' not in parsed['extracted_text']
    pdf=PdfWriter();pdf.add_blank_page(width=200,height=200);path=tmp_path/'scan.pdf'
    with path.open('wb') as stream:pdf.write(stream)
    parsed=inspect_file(path,'scan.pdf');assert parsed['text_status']=='unsupported' and parsed['extracted_text']==''
    path=tmp_path/'long.txt';path.write_text('字'*200001)
    parsed=inspect_file(path,'long.txt');assert parsed['text_status']=='truncated' and len(parsed['extracted_text'])==200000


def test_zip_bomb_and_private_path_symlink_rejected(tmp_path):
    import zipfile
    from services.knowledge_workspace.files import inspect_file,private_path
    from core.exceptions import AppException
    archive=tmp_path/'bomb.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('large.txt',b'x'*(50*1024*1024+1))
    with pytest.raises(AppException,match='压缩内容'):inspect_file(archive,'bomb.zip')
    root=tmp_path/'nas';root.mkdir();(root/'.knowledge-versions').symlink_to(tmp_path,target_is_directory=True)
    with pytest.raises(AppException,match='存储不可用'):private_path(root,None,uuid4(),uuid4(),create=True)


def test_oversized_upload_and_cancelled_receive_leave_no_partial_file(file_client,tmp_path):
    import asyncio
    from services.workspace_coordination import receive_workspace_upload
    client,_,settings,r,_=file_client;obj=document(r)
    response=client.post(f'/api/knowledge/objects/{obj["id"]}/files',data={'command_id':str(uuid4())},files={'file':('big.txt',b'x'*(20*1024*1024+1))})
    assert response.status_code==413
    assert client.get(f'/api/knowledge/objects/{obj["id"]}/files').json()['data']==[]
    assert not list(Path(settings.file_workspace_root).rglob('.upload-*'))
    class Interrupted:
        calls=0
        async def read(self,_):
            self.calls+=1
            if self.calls==1:return b'partial'
            raise asyncio.CancelledError()
    async def receive():
        with pytest.raises(asyncio.CancelledError):
            await receive_workspace_upload(Interrupted(),tmp_path/'partial',max_bytes=100,too_large=ValueError)
    asyncio.run(receive());assert not (tmp_path/'partial').exists() and not list(tmp_path.glob('.upload-*'))


def test_unauthenticated_private_file_request_is_rejected_without_cache(file_client):
    from api.routes import knowledge_files
    from core.database import get_db
    client,ctx,settings,r,_=file_client
    app=FastAPI();app.include_router(knowledge_files.router,prefix='/api')
    app.dependency_overrides[get_db]=lambda:SimpleNamespace(pool=r.pool)
    with TestClient(app,raise_server_exceptions=False) as unauthenticated:
        response=unauthenticated.get(f'/api/knowledge/objects/{uuid4()}/files')
        assert response.status_code==401 and response.headers['cache-control']=='no-store'


def test_kill_switch_blocks_all_private_file_reads_and_writes(file_client):
    client,_,settings,r,_=file_client;obj=document(r);bid=uploaded(client,obj['id'])['binding_id']
    settings.knowledge_workspace_enabled=False
    assert client.get(url(obj['id'],bid)).status_code==404
    assert client.get(f'/api/knowledge/objects/{obj["id"]}/files').status_code==404
    response=client.post(f'/api/knowledge/objects/{obj["id"]}/files',data={'command_id':str(uuid4())},files={'file':('no.txt',b'no')})
    assert response.status_code==404


def test_generic_asset_gallery_excludes_private_knowledge_names_counts_and_urls(file_client):
    from psycopg.types.json import Jsonb
    client,ctx,_,r,make=file_client;obj=document(r);uploaded(client,obj['id'])
    private=client.get(f'/api/knowledge/objects/{obj["id"]}/files').json()['data'][0]['version']['asset_id']
    with psycopg.connect(**make.admin_args) as c:
        ordinary=uuid4()
        c.execute("""INSERT INTO user_assets(id,org_id,storage_scope,storage_owner_key,storage_provider,storage_key,media_type,original_url,download_url,name)
          VALUES(%s,%s,'user',%s,'oss',%s,'file','https://example.test/ordinary','https://example.test/ordinary','普通上传.txt')""",(ordinary,ctx.org_id,ctx.user_id,str(ordinary)))
        c.execute("""INSERT INTO user_asset_refs(ref_key,asset_id,actor_user_id,org_id,source_type,source_kind,ref_kind)
          VALUES(%s,%s,%s,%s,'upload','web_upload','upload')""",(str(ordinary),ordinary,ctx.user_id,ctx.org_id))
        for name in ['list_admin_user_assets','list_platform_admin_user_assets']:
            result=c.execute(f'SELECT {name}(%s,%s)',(ctx.user_id,'upload')).fetchone()[0]
            assert result['total']==1 and [row['name'] for row in result['items']]==['普通上传.txt']
        with pytest.raises(psycopg.Error,match='ADMIN_ASSET_DOWNLOAD_SCOPE_INVALID'):
            c.execute('SELECT resolve_platform_admin_user_assets_download(%s,%s)',(ctx.user_id,Jsonb([private])))
        c.rollback()


def test_concurrent_file_updates_do_not_silently_overwrite(file_client):
    from concurrent.futures import ThreadPoolExecutor
    from services.knowledge_workspace.files import version_id
    client,ctx,_,r,_=file_client;obj=document(r);bid=uploaded(client,obj['id'])['binding_id']
    def update(index):
        key=uuid4()
        try:
            return r.file_command(key,'file_upload',{'object_id':str(obj['id']),'binding_id':bid,'expected_revision':1,
                'version_id':version_id(ctx.user_id,ctx.org_id,key),'name':f'writer-{index}.txt','mime_type':'text/plain',
                'size':1,'sha256':str(index)*64,'text_status':'unsupported'})['status']
        except psycopg.Error as error:return error.diag.message_primary
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(update,[1,2]))
    assert sorted(results)==['KB_VERSION_CONFLICT','applied']
    versions=r.file_versions(obj['id'],bid)['items'];assert [v['revision'] for v in versions]==[2,1]
