"""Bounded CSV/XLSX preview then atomic confirmation; no silent replacement."""

import csv
import io
from itertools import islice
from datetime import datetime, timezone
from uuid import UUID
from psycopg.types.json import Jsonb
from services.agent.agent_reach.contracts import ReachError
from .analytics import METRICS, exact_metric, timestamp
from .metrics import verified, save_snapshot
from .repository import digest

FIELDS = (
    "intent_id",
    "account_id",
    "remote_id",
    "observed_at",
    "published_at",
    "content_type",
    "scope",
    "denominator",
    *METRICS,
)
REQUIRED = (
    "intent_id",
    "account_id",
    "remote_id",
    "observed_at",
    "content_type",
    "scope",
)


def parse_table(raw, filename):
    if len(raw) > 2 * 1024 * 1024:
        raise ReachError("IMPORT_LIMIT", "文件最多2MB、500行、30列")
    if filename.lower().endswith(".csv"):
        try:
            table = list(
                islice(
                    csv.reader(io.StringIO(raw.decode("utf-8-sig")), strict=True), 502
                )
            )
        except (UnicodeError, csv.Error):
            raise ReachError("INVALID_ARGUMENTS", "CSV必须为UTF-8且格式有效") from None
    elif filename.lower().endswith(".xlsx"):
        from services.skills.imports import _check_office
        from services.skills.contracts import SkillError
        from openpyxl import load_workbook

        try:
            _check_office(raw)
            wb = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(wb.worksheets) != 1:
                    raise ReachError(
                        "INVALID_ARGUMENTS", "Excel仅支持一个工作表，请按模板导出"
                    )
                sheet = wb.worksheets[0]
                sheet.reset_dimensions()
                table = [
                    list(r)
                    for r in sheet.iter_rows(max_row=502, max_col=31, values_only=True)
                ]
                table = [[str(v) if v is not None else "" for v in r] for r in table]
                table = [r for r in table if any(v for v in r)]
                table = [
                    r[: max((i + 1 for i, v in enumerate(r) if v), default=0)]
                    for r in table
                ]
            finally:
                wb.close()
        except ReachError:
            raise
        except Exception:
            raise ReachError(
                "INVALID_ARGUMENTS", "Excel无法解析、包含宏/嵌入对象或超出复杂度限制"
            ) from None
    else:
        raise ReachError("INVALID_ARGUMENTS", "仅支持.csv或.xlsx，不支持宏和旧版xls")
    if (
        not table
        or len(table) < 2
        or len(table) > 501
        or any(len(r) > 30 for r in table)
    ):
        raise ReachError("IMPORT_LIMIT", "需要表头和1至500行数据、最多30列")
    headers = [str(v).strip() for v in table[0]]
    if len(set(headers)) != len(headers) or any(not h for h in headers):
        raise ReachError("INVALID_ARGUMENTS", "表头不能为空或重复")
    rows = []
    for number, row in enumerate(table[1:], 2):
        if len(row) > len(headers) or any(
            len(str(v)) > 4000 or str(v).lstrip().startswith(("=", "+", "@"))
            for v in row
        ):
            raise ReachError(
                "INVALID_ARGUMENTS", f"第{number}行列数错误、内容过长或包含公式"
            )
        rows.append(
            dict(
                zip(
                    headers,
                    [str(v).strip() for v in row] + [""] * (len(headers) - len(row)),
                )
            )
        )
    return headers, rows


def preview(repo, raw, filename, source, mapping):
    headers, rows = parse_table(raw, filename)
    mapping = mapping or {k: k for k in FIELDS if k in headers}
    if (
        set(mapping) - set(FIELDS)
        or any(v not in headers for v in mapping.values())
        or len(set(mapping.values())) != len(mapping)
    ):
        raise ReachError("INVALID_ARGUMENTS", "字段映射无效或重复")
    if not set(REQUIRED) <= set(mapping):
        raise ReachError("INVALID_ARGUMENTS", "必填字段：" + ",".join(REQUIRED))
    output = []
    seen = set()
    with repo.cursor() as c:
        for number, raw_row in enumerate(rows, 2):
            row = {k: raw_row.get(mapping.get(k, ""), "") for k in FIELDS}
            errors = []
            intent = None
            try:
                intent = verified(repo, c, str(UUID(row["intent_id"])))
                if (
                    intent["account_id"] != row["account_id"]
                    or intent["verification"]["remote_id"] != row["remote_id"]
                ):
                    errors.append("账号或平台内容ID与核验记录不符")
            except (ValueError, ReachError):
                errors.append("发布内容不可访问或尚未核验")
            observed = timestamp(row["observed_at"])
            published = timestamp(row["published_at"])
            if observed is None or observed > datetime.now(timezone.utc):
                errors.append("采集时间需有效时区且不能在未来")
            if row["published_at"] and (
                not published or observed and published > observed
            ):
                errors.append("发布时间无效或晚于采集")
            if row["content_type"] not in {"image", "video", "text"}:
                errors.append("内容类型应为image/video/text")
            if row["scope"] not in {"creator_cumulative", "public_cumulative"}:
                errors.append("scope应为creator_cumulative/public_cumulative")
            if row["denominator"] not in {"", "views", "impressions"}:
                errors.append("denominator应为空/views/impressions")
            metrics = {}
            for k in METRICS:
                metrics[k] = None if row[k] == "" else exact_metric(k, row[k])
                if row[k] and metrics[k] is None:
                    errors.append(
                        k
                        + (
                            "需为空或0至1的原始比例，不带百分号"
                            if k == "completion_rate"
                            else "必须为空或精确非负整数"
                        )
                    )
            if not any(v is not None for v in metrics.values()):
                errors.append("至少提供一个真实指标")
            normalized = {
                "intent_id": row["intent_id"],
                "account_id": row["account_id"],
                "remote_id": row["remote_id"],
                "observed_at": observed.isoformat() if observed else None,
                "published_at": published.isoformat() if published else None,
                "content_type": row["content_type"],
                "metrics": metrics,
                "coverage": {
                    "scope": row["scope"],
                    "denominator": row["denominator"] or None,
                    "fields": [k for k, v in metrics.items() if v is not None],
                    "publication_source": "import" if published else "unknown",
                },
                "source": source,
                "adapter_version": "controlled-import:v1",
            }
            # Semantic identity ignores file ordering/name/importer. Same timestamp with
            # changed values is a conflict, never a replacement or duplicate snapshot.
            key = "import:" + digest(
                {
                    k: normalized[k]
                    for k in ("intent_id", "account_id", "remote_id", "observed_at")
                }
            )
            duplicate = False
            if not errors:
                c.execute(
                    "SELECT metrics,published_at,coverage,content_type FROM social_metric_snapshots WHERE org_id=%s AND dedupe_key=%s",
                    (repo.scope.org_id, key),
                )
                existing = c.fetchone()
                if existing:
                    if (
                        existing["metrics"] != metrics
                        or timestamp(existing["published_at"]) != published
                        or existing["coverage"] != normalized["coverage"]
                        or existing["content_type"] != row["content_type"]
                    ):
                        errors.append(
                            "同一内容和采集时间已有不同数据，禁止静默覆盖；请核实原始来源"
                        )
                    else:
                        duplicate = True
            if key in seen:
                errors.append("文件内同一内容和采集时间重复")
            seen.add(key)
            output.append(
                {
                    "line": number,
                    "data": normalized,
                    "dedupe_key": key,
                    "errors": errors,
                    "duplicate": duplicate,
                }
            )
        value = {
            "rows": output,
            "valid": not any(r["errors"] for r in output),
            "duplicate_count": sum(r["duplicate"] for r in output),
            "rules": "全部有效后原子确认；相同内容/采集时间/数据复用，冲突拒绝；不补零，不覆盖",
        }
        c.execute(
            """INSERT INTO social_metric_imports(org_id,owner_id,file_hash,source,mapping,preview,preview_hash)
            VALUES(%s,%s,%s,%s,%s,%s,%s) RETURNING id,preview_hash""",
            (
                repo.scope.org_id,
                repo.scope.actor_user_id,
                digest(raw.hex()),
                source,
                Jsonb(mapping),
                Jsonb(value),
                digest(value),
            ),
        )
        return {**c.fetchone(), **value}


def confirm(repo, import_id, preview_hash):
    with repo.cursor() as c:
        c.execute(
            "SELECT * FROM social_metric_imports WHERE id=%s AND org_id=%s AND owner_id=%s FOR UPDATE",
            (import_id, repo.scope.org_id, repo.scope.actor_user_id),
        )
        batch = c.fetchone()
        if not batch:
            raise ReachError("NOT_FOUND", "导入预览不存在或不属于当前操作者")
        if batch["preview_hash"] != preview_hash or not batch["preview"]["valid"]:
            raise ReachError("VERSION_CONFLICT", "预览无效或已经变化")
        if batch["status"] == "confirmed":
            return {
                "status": "already_confirmed",
                "created": 0,
                "reused": len(batch["preview"]["rows"]),
            }
        if (datetime.now(timezone.utc) - batch["created_at"]).total_seconds() > 3600:
            raise ReachError("VERSION_CONFLICT", "预览已过期，请重新预览")
        created = reused = 0
        for item in batch["preview"]["rows"]:
            row = item["data"]
            intent = verified(repo, c, row["intent_id"])
            if (
                intent["account_id"] != row["account_id"]
                or intent["verification"]["remote_id"] != row["remote_id"]
            ):
                raise ReachError("VERSION_CONFLICT", "内容核验关联已变化，请重新预览")
            saved, duplicate = save_snapshot(
                repo,
                c,
                intent,
                {**row, "coverage": row["coverage"]},
                item["dedupe_key"],
            )
            if (
                saved["metrics"] != row["metrics"]
                or timestamp(saved["published_at"]) != timestamp(row["published_at"])
                or saved["coverage"] != row["coverage"]
                or saved["content_type"] != row["content_type"]
            ):
                raise ReachError("VERSION_CONFLICT", "并发导入数据冲突，整个批次未写入")
            reused += int(duplicate)
            created += int(not duplicate)
        c.execute(
            "UPDATE social_metric_imports SET status='confirmed',confirmed_at=now() WHERE id=%s",
            (import_id,),
        )
        return {"status": "confirmed", "created": created, "reused": reused}
