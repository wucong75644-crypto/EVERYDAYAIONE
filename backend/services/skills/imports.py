"""Import document bytes as reviewed text plus an immutable original file.

Only fixed library parsers run here. No uploaded code, macros, shell commands,
workspace paths, URL fetches or model/tool execution are involved.
"""

from io import BytesIO
import multiprocessing
from pathlib import PurePosixPath
import re
import sys
import threading
from uuid import uuid4
from zipfile import ZipFile
from xml.etree import ElementTree

from services.skills.assets import AssetDraft, AssetSourceDraft, MAX_ASSET_BYTES, MAX_SOURCE_BYTES
from services.skills.contracts import SkillError

FORMATS = frozenset({'txt', 'md', 'csv', 'json', 'docx', 'pdf', 'xlsx'})
_PARSERS = threading.BoundedSemaphore(2)
PARSE_TIMEOUT = 10


class _Text:
    def __init__(self):
        self.parts = []
        self.bytes = 0

    def add(self, value):
        if not value or not value.strip():
            return
        self.bytes += len(value.encode('utf-8')) + bool(self.parts)
        if self.bytes > MAX_ASSET_BYTES:
            raise SkillError('SKILL_UPLOAD_TEXT_TOO_LARGE')
        if '\x00' in value:
            raise SkillError('SKILL_UPLOAD_INVALID')
        self.parts.append(value)

    def finish(self):
        if not self.parts:
            raise SkillError('SKILL_UPLOAD_NO_TEXT')
        return '\n'.join(self.parts)


def _check_office(raw):
    # Inspect compressed containers before a parser inflates their XML. Never
    # extract entries onto disk; external relationships are never followed.
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 512 or sum(e.file_size for e in entries) > 16 * 1024 * 1024:
            raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT')
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if (path.is_absolute() or '..' in path.parts or '\\' in entry.filename
                    or entry.flag_bits & 1 or entry.file_size > 8 * 1024 * 1024
                    or entry.filename.lower().endswith('vbaproject.bin')
                    or '/embeddings/' in entry.filename.lower()):
                raise SkillError('SKILL_UPLOAD_INVALID')
            if entry.filename.endswith(('.xml', '.rels')):
                xml = archive.read(entry)
                if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                    raise SkillError('SKILL_UPLOAD_INVALID')
                if entry.filename.startswith('xl/worksheets/'):
                    for event, cell in ElementTree.iterparse(BytesIO(xml), events=('start',)):
                        if cell.tag.rsplit('}', 1)[-1] != 'c':
                            continue
                        address = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)', cell.get('r', ''))
                        if not address:
                            raise SkillError('SKILL_UPLOAD_INVALID')
                        column = 0
                        for letter in address[1]:
                            column = column * 26 + ord(letter) - ord('A') + 1
                        if column > 100 or int(address[2]) > 2000:
                            raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT')


def _extract(format, raw):
    text = _Text()
    if format in ('txt', 'md', 'csv', 'json'):
        try:
            value = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise SkillError('SKILL_UPLOAD_ENCODING') from None
        text.add(value)
    elif format == 'docx':
        _check_office(raw)
        from docx import Document
        from docx.table import Table
        for block in Document(BytesIO(raw)).iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    text.add('\t'.join(cell.text for cell in row.cells))
            else:
                text.add(block.text)
    elif format == 'xlsx':
        _check_office(raw)
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > 16:
                raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT')
            for sheet in workbook:
                # Do not trust dimensions provided by the uploaded XML.
                sheet.reset_dimensions()
                named = False
                for number, row in enumerate(sheet.iter_rows(max_row=2001, max_col=101, values_only=True), 1):
                    if not any(cell is not None for cell in row):
                        continue
                    if number > 2000 or row[100] is not None:
                        raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT')
                    if not named:
                        text.add(f'[{sheet.title}]')
                        named = True
                    text.add('\t'.join('' if cell is None else str(cell) for cell in row[:100]).rstrip('\t'))
        finally:
            workbook.close()
    elif format == 'pdf':
        from pypdf import PdfReader
        document = PdfReader(BytesIO(raw), strict=True)
        if document.is_encrypted:
            raise SkillError('SKILL_UPLOAD_ENCRYPTED')
        if len(document.pages) > 100:
            raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT')
        for page in document.pages:
            text.add(page.extract_text())
    else:
        raise SkillError('SKILL_UPLOAD_FORMAT')
    return text.finish()


def _parse_worker(connection, format, raw):
    try:
        # A small compressed PDF can inflate before extract_text returns. A
        # disposable parser process bounds that work without blocking the API.
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (5, 6))
        if sys.platform == 'linux':
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024,) * 2)
        connection.send((True, _extract(format, raw)))
    except SkillError as error:
        connection.send((False, str(error)))
    except Exception:
        connection.send((False, 'SKILL_UPLOAD_INVALID'))
    finally:
        connection.close()


def _parse_document(format, raw):
    if not _PARSERS.acquire(blocking=False):
        raise SkillError('SKILL_UPLOAD_BUSY')
    context = multiprocessing.get_context('spawn')
    reader, writer = context.Pipe(duplex=False)
    process = context.Process(target=_parse_worker, args=(writer, format, raw), daemon=True)
    try:
        process.start()
        writer.close()
        if not reader.poll(PARSE_TIMEOUT):
            raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT')
        try:
            ok, result = reader.recv()
        except EOFError:
            raise SkillError('SKILL_UPLOAD_COMPLEXITY_LIMIT') from None
        if not ok:
            raise SkillError(result)
        return result
    finally:
        if process.pid is not None:
            process.join(timeout=0.2)
            if process.is_alive():
                process.kill()
                process.join(timeout=1)
        reader.close()
        writer.close()
        _PARSERS.release()


def import_attachment(filename: str, raw: bytes) -> AssetDraft:
    if (not filename or len(filename) > 120 or any(c in filename for c in '/\\')
            or any(ord(c) < 32 for c in filename)):
        raise SkillError('SKILL_UPLOAD_FILENAME_INVALID')
    format = PurePosixPath(filename).suffix[1:].lower()
    if format not in FORMATS:
        raise SkillError('SKILL_UPLOAD_FORMAT')
    if len(raw) > MAX_SOURCE_BYTES:
        raise SkillError('SKILL_UPLOAD_TOO_LARGE')
    if not raw:
        raise SkillError('SKILL_UPLOAD_NO_TEXT')
    content = (_extract(format, raw) if format in ('txt', 'md', 'csv', 'json')
               else _parse_document(format, raw))
    return AssetDraft(id='attachment-' + uuid4().hex, name=filename, kind='reference',
        summary=f'参考资料：{filename}', format=format if format in ('md', 'txt', 'json', 'csv') else 'txt',
        content=content, source=AssetSourceDraft.from_bytes(format, raw))
