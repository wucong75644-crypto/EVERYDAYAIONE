"""Bounded extraction worker. Invoked as a subprocess, without application imports."""
import json
from pathlib import Path
import sys

LIMIT = 200000


def extract(path, extension):
    if extension == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('encrypted')
        chunks = []
        total = 0
        for page in reader.pages[:200]:
            value = page.extract_text() or ''
            chunks.append(value[:LIMIT + 1 - total]); total += len(chunks[-1])
            if total > LIMIT: break
        return '\n'.join(chunks), len(reader.pages) > 200
    if extension == '.docx':
        from docx import Document
        document = Document(path)
        chunks = []
        total = 0
        for paragraph in document.paragraphs:
            chunks.append(paragraph.text[:LIMIT + 1 - total]); total += len(chunks[-1]) + 1
            if total > LIMIT: return '\n'.join(chunks), True
        for table in document.tables:
            for row in table.rows:
                value = '\t'.join(cell.text for cell in row.cells)
                chunks.append(value[:LIMIT + 1 - total]); total += len(chunks[-1]) + 1
                if total > LIMIT: return '\n'.join(chunks), True
        return '\n'.join(chunks), False
    if extension == '.xlsx':
        from openpyxl import load_workbook
        stream = open(path,'rb')
        book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        chunks = []; total = 0; partial = False
        try:
            for sheet in book.worksheets:
                partial = partial or (sheet.max_column or 0)>200
                for index, row in enumerate(sheet.iter_rows(max_col=200, values_only=True)):
                    if index >= 10000: return '\n'.join(chunks), True
                    value = '\t'.join(str(v) if v is not None else '' for v in row)
                    chunks.append(value[:LIMIT + 1 - total]); total += len(chunks[-1]) + 1
                    if total > LIMIT: return '\n'.join(chunks), True
            return '\n'.join(chunks), partial
        finally:
            book.close(); stream.close()
    return '', False


if __name__ == '__main__':
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        # Linux supports AS; macOS does not reliably enforce it.
        if sys.platform.startswith('linux'):
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        value, partial = extract(sys.argv[1], sys.argv[2])
        value = value.replace('\x00', '')
        print(json.dumps({'extracted_text': value[:LIMIT], 'text_status': 'unsupported' if not value.strip() else ('truncated' if partial or len(value)>LIMIT else 'ready')}))
    except Exception:
        print(json.dumps({'extracted_text': '', 'text_status': 'failed'}))
